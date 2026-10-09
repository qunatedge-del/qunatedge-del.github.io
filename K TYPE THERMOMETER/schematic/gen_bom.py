# -*- coding: utf-8 -*-
"""Writes ../bom/ktype_thermometer_BOM.xlsx and .csv from design_data.PARTS"""
import csv
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from design_data import PARTS, CATEGORY_ORDER, PROJECT, REV, DATE

OUT = os.path.join(HERE, "..", "bom")
os.makedirs(OUT, exist_ok=True)


def refkey(r):
    m = re.match(r"([A-Za-z]+)(\d+)", r)
    return (m.group(1), int(m.group(2)))


groups = {}
for ref, p in PARTS.items():
    key = (p["cat"], p["value"], p["pkg"], p["desc"], p["mfr"], p["mpn"], p["notes"])
    groups.setdefault(key, []).append(ref)

rows = []
for key, refs in groups.items():
    refs.sort(key=refkey)
    rows.append((CATEGORY_ORDER.index(key[0]), refkey(refs[0]), key, refs))
rows.sort(key=lambda t: (t[0], t[1]))

HEAD = ["Item #", "Ref Des", "Qty", "Value", "Package", "Description", "Manufacturer", "MPN", "Category", "Notes"]
table = []
n = 0
for _, _, key, refs in rows:
    n += 1
    cat, val, pkg, desc, mfr, mpn, notes = key
    table.append([n, ", ".join(refs), len(refs), val, pkg, desc, mfr, mpn, cat, notes])

with open(os.path.join(OUT, "ktype_thermometer_BOM.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(HEAD)
    w.writerows(table)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "BOM"
ws.append(["%s  -  BOM  (Rev %s, %s)" % (PROJECT, REV, DATE)])
ws["A1"].font = Font(bold=True, size=14)
ws.append(["Passives are generic (spec in Value/Notes). ICs carry the suggested MPN. DNP = do not populate."])
ws["A2"].font = Font(italic=True, color="666666")
ws.append([])
ws.append(HEAD)
hr = 4
thin = Side(style="thin", color="BBBBBB")
for c in range(1, len(HEAD) + 1):
    cell = ws.cell(row=hr, column=c)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="37474F")
    cell.alignment = Alignment(vertical="center", horizontal="center")

cat_fill = {"IC": "E8EAF6", "Discrete": "FFF3E0", "Passive": "F1F8E9", "Crystal": "FFF8E1",
            "Connector": "E3F2FD", "Electromech": "F3E5F5"}
row = hr + 1
first_row = {}
last_row = {}
for item in table:
    ws.append(item)
    cat = item[8]
    first_row.setdefault(cat, row)
    last_row[cat] = row
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=cat_fill[cat])
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    row += 1

ws.append([])
row += 1
ws.cell(row=row, column=2, value="Subtotals by category (quantity of parts)").font = Font(bold=True)
row += 1
sub_rows = []
for cat in CATEGORY_ORDER:
    if cat in first_row:
        ws.cell(row=row, column=2, value=cat)
        ws.cell(row=row, column=3, value="=SUM(C%d:C%d)" % (first_row[cat], last_row[cat]))
        # only rows of this category are contiguous because the list is sorted by category
        sub_rows.append(row)
        row += 1
ws.cell(row=row, column=2, value="TOTAL parts").font = Font(bold=True)
ws.cell(row=row, column=3, value="=SUM(C%d:C%d)" % (sub_rows[0], sub_rows[-1])).font = Font(bold=True)

widths = [7, 34, 6, 22, 18, 52, 24, 24, 13, 60]
for i, w_ in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w_
ws.freeze_panes = "A5"
ws.sheet_view.zoomScale = 90
wb.save(os.path.join(OUT, "ktype_thermometer_BOM.xlsx"))
total = sum(len(r[3]) for r in rows)
print("BOM lines:", len(table), " parts:", total, " (PARTS:", len(PARTS), ")")
assert total == len(PARTS)
